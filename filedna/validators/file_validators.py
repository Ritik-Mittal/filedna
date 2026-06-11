"""
FileDNA – validators.

Each validator returns a ValidatorResult (is_valid, errors, warnings).
Validators use magic bytes + library-level structural checks.

Design notes:
- ValidationContext dataclass eliminates boilerplate (no more repeated list init)
- All imports are lazy (inside functions) for minimal startup cost
- _DISPATCH table maps type string → validator function
"""
from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path

# ---------------------------------------------------------------------------
# Result types
# ---------------------------------------------------------------------------

@dataclass
class ValidationContext:
    """Accumulates errors and warnings during validation. Eliminates boilerplate."""
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def fail(self, msg: str) -> ValidatorResult:
        self.errors.append(msg)
        return False, self.errors, self.warnings

    def warn(self, msg: str) -> None:
        self.warnings.append(msg)

    def ok(self) -> ValidatorResult:
        return True, self.errors, self.warnings


ValidatorResult = tuple[bool, list[str], list[str]]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _read_bytes(path: Path, n: int = 16) -> bytes:
    try:
        with open(path, "rb") as f:
            return f.read(n)
    except OSError:
        return b""


# ---------------------------------------------------------------------------
# Document validators
# ---------------------------------------------------------------------------

def validate_pdf(path: Path) -> ValidatorResult:
    ctx = ValidationContext()
    head = _read_bytes(path, 8)
    if not head.startswith(b"%PDF-"):
        return ctx.fail("File is not a valid PDF")
    try:
        import pdfplumber  # type: ignore
        with pdfplumber.open(str(path)) as pdf:
            if len(pdf.pages) == 0:
                ctx.warn("PDF has no pages")
    except Exception as exc:
        return ctx.fail(f"PDF is corrupted or unreadable: {exc}")
    return ctx.ok()


def validate_docx(path: Path) -> ValidatorResult:
    ctx = ValidationContext()
    import zipfile
    try:
        with zipfile.ZipFile(path) as z:
            if "word/document.xml" not in z.namelist():
                return ctx.fail("Missing word/document.xml – not a valid DOCX")
    except zipfile.BadZipFile:
        return ctx.fail("File is not a valid DOCX (bad ZIP structure)")
    except Exception as exc:
        return ctx.fail(f"DOCX validation failed: {exc}")
    try:
        from docx import Document  # type: ignore
        Document(str(path))
    except Exception as exc:
        return ctx.fail(f"DOCX is corrupted: {exc}")
    return ctx.ok()


def validate_xlsx(path: Path) -> ValidatorResult:
    ctx = ValidationContext()
    import zipfile
    try:
        with zipfile.ZipFile(path) as z:
            if "xl/workbook.xml" not in z.namelist():
                return ctx.fail("Missing xl/workbook.xml – not a valid XLSX")
    except zipfile.BadZipFile:
        return ctx.fail("File is not a valid XLSX (bad ZIP structure)")
    except Exception as exc:
        return ctx.fail(f"XLSX validation failed: {exc}")
    try:
        import openpyxl  # type: ignore
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        wb.close()
    except Exception as exc:
        return ctx.fail(f"XLSX is corrupted: {exc}")
    return ctx.ok()


def validate_pptx(path: Path) -> ValidatorResult:
    ctx = ValidationContext()
    import zipfile
    try:
        with zipfile.ZipFile(path) as z:
            if "ppt/presentation.xml" not in z.namelist():
                return ctx.fail("Missing ppt/presentation.xml – not a valid PPTX")
    except zipfile.BadZipFile:
        return ctx.fail("File is not a valid PPTX (bad ZIP structure)")
    except Exception as exc:
        return ctx.fail(f"PPTX validation failed: {exc}")
    try:
        from pptx import Presentation  # type: ignore
        Presentation(str(path))
    except Exception as exc:
        return ctx.fail(f"PPTX is corrupted: {exc}")
    return ctx.ok()


def validate_epub(path: Path) -> ValidatorResult:
    ctx = ValidationContext()
    import zipfile
    try:
        with zipfile.ZipFile(path) as z:
            if "META-INF/container.xml" not in z.namelist():
                return ctx.fail("Missing META-INF/container.xml – not a valid EPUB")
    except zipfile.BadZipFile:
        return ctx.fail("File is not a valid EPUB (bad ZIP structure)")
    except Exception as exc:
        return ctx.fail(f"EPUB validation failed: {exc}")
    return ctx.ok()


def validate_csv(path: Path) -> ValidatorResult:
    ctx = ValidationContext()
    try:
        import csv
        with open(path, newline="", encoding="utf-8", errors="replace") as f:
            rows = list(csv.reader(f))
        if not rows:
            ctx.warn("CSV file is empty")
    except Exception as exc:
        return ctx.fail(f"CSV is unreadable: {exc}")
    return ctx.ok()


def _validate_text(path: Path, type_name: str) -> ValidatorResult:
    ctx = ValidationContext()
    try:
        content = path.read_text(encoding="utf-8", errors="replace")
        if not content.strip():
            ctx.warn(f"{type_name.upper()} file is empty")
    except Exception as exc:
        return ctx.fail(f"{type_name.upper()} is unreadable: {exc}")
    return ctx.ok()


def validate_txt(path: Path) -> ValidatorResult:
    return _validate_text(path, "txt")

def validate_md(path: Path) -> ValidatorResult:
    return _validate_text(path, "md")

def validate_html(path: Path) -> ValidatorResult:
    return _validate_text(path, "html")


def validate_xml(path: Path) -> ValidatorResult:
    ctx = ValidationContext()
    try:
        import xml.etree.ElementTree as ET
        ET.parse(str(path))
    except Exception as exc:
        return ctx.fail(f"XML parse error: {exc}")
    return ctx.ok()


def validate_json(path: Path) -> ValidatorResult:
    ctx = ValidationContext()
    try:
        import json
        with open(path, encoding="utf-8", errors="replace") as f:
            json.load(f)
    except Exception as exc:
        return ctx.fail(f"JSON parse error: {exc}")
    return ctx.ok()


# ---------------------------------------------------------------------------
# Image validators
# ---------------------------------------------------------------------------

def validate_image(path: Path) -> ValidatorResult:
    ctx = ValidationContext()
    try:
        from PIL import Image  # type: ignore
        with Image.open(str(path)) as img:
            img.verify()
    except Exception as exc:
        return ctx.fail(f"Image is corrupted or unreadable: {exc}")
    return ctx.ok()


def validate_svg(path: Path) -> ValidatorResult:
    return validate_xml(path)


# ---------------------------------------------------------------------------
# Audio validators
# ---------------------------------------------------------------------------

def validate_audio(path: Path) -> ValidatorResult:
    ctx = ValidationContext()
    try:
        import mutagen  # type: ignore
        tag = mutagen.File(str(path))
        if tag is None:
            ctx.warn("Audio metadata could not be read")
    except Exception as exc:
        return ctx.fail(f"Audio file is unreadable: {exc}")
    return ctx.ok()


# ---------------------------------------------------------------------------
# Video validators
# ---------------------------------------------------------------------------

def validate_video(path: Path) -> ValidatorResult:
    ctx = ValidationContext()
    head = _read_bytes(path, 16)
    if head[4:8] == b"ftyp":       # mp4 / mov
        return ctx.ok()
    if head[:4] == b"\x1a\x45\xdf\xa3":  # mkv / webm
        return ctx.ok()
    if head[:4] == b"RIFF" and head[8:12] == b"AVI ":
        return ctx.ok()
    ctx.warn("Could not fully verify video structure; file may be valid")
    return ctx.ok()


# ---------------------------------------------------------------------------
# Archive validators
# ---------------------------------------------------------------------------

def validate_zip(path: Path) -> ValidatorResult:
    ctx = ValidationContext()
    import zipfile
    try:
        with zipfile.ZipFile(path) as z:
            bad = z.testzip()
            if bad:
                return ctx.fail(f"ZIP integrity check failed on: {bad}")
    except zipfile.BadZipFile as exc:
        return ctx.fail(f"Invalid ZIP file: {exc}")
    except Exception as exc:
        return ctx.fail(f"ZIP validation error: {exc}")
    return ctx.ok()


def validate_tar(path: Path) -> ValidatorResult:
    ctx = ValidationContext()
    import tarfile
    try:
        with tarfile.open(str(path)) as t:
            t.getmembers()
    except Exception as exc:
        return ctx.fail(f"TAR validation error: {exc}")
    return ctx.ok()


def validate_gz(path: Path) -> ValidatorResult:
    ctx = ValidationContext()
    import gzip
    try:
        with gzip.open(str(path), "rb") as f:
            f.read(1024)
    except Exception as exc:
        return ctx.fail(f"GZ validation error: {exc}")
    return ctx.ok()


def validate_bz2(path: Path) -> ValidatorResult:
    ctx = ValidationContext()
    import bz2
    try:
        with bz2.open(str(path), "rb") as f:
            f.read(1024)
    except Exception as exc:
        return ctx.fail(f"BZ2 validation error: {exc}")
    return ctx.ok()


def validate_7z(path: Path) -> ValidatorResult:
    ctx = ValidationContext()
    if _read_bytes(path, 6) != b"7z\xbc\xaf'\x1c":
        return ctx.fail("Invalid 7Z magic bytes")
    return ctx.ok()


def validate_rar(path: Path) -> ValidatorResult:
    ctx = ValidationContext()
    if not _read_bytes(path, 7).startswith(b"Rar!\x1a\x07"):
        return ctx.fail("Invalid RAR magic bytes")
    return ctx.ok()


# ---------------------------------------------------------------------------
# Dispatch table
# ---------------------------------------------------------------------------

_IMAGE_TYPES = {"png", "jpg", "jpeg", "gif", "bmp", "webp", "tiff", "tif"}
_AUDIO_TYPES = {"mp3", "wav", "flac", "ogg", "m4a", "aac"}
_VIDEO_TYPES = {"mp4", "mov", "mkv", "webm", "avi"}

_DISPATCH: dict[str, Callable[[Path], ValidatorResult]] = {
    "pdf":  validate_pdf,
    "docx": validate_docx,
    "xlsx": validate_xlsx,
    "pptx": validate_pptx,
    "epub": validate_epub,
    "csv":  validate_csv,
    "txt":  validate_txt,
    "md":   validate_md,
    "html": validate_html,
    "xml":  validate_xml,
    "json": validate_json,
    "svg":  validate_svg,
    "zip":  validate_zip,
    "tar":  validate_tar,
    "gz":   validate_gz,
    "bz2":  validate_bz2,
    "7z":   validate_7z,
    "rar":  validate_rar,
    **dict.fromkeys(_IMAGE_TYPES, validate_image),
    **dict.fromkeys(_AUDIO_TYPES, validate_audio),
    **dict.fromkeys(_VIDEO_TYPES, validate_video),
}


def validate(path: Path, real_type: str) -> ValidatorResult:
    """Validate *path* as *real_type*. Returns (valid, errors, warnings)."""
    fn = _DISPATCH.get(real_type)
    if fn is None:
        return True, [], [f"No validator for type '{real_type}'; skipping structural check"]
    return fn(path)
